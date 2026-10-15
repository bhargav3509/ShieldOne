import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_shieldone_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide

    # Color Palette
    COLOR_BG = RGBColor(255, 255, 255)
    COLOR_NAVY = RGBColor(15, 23, 42)       # #0F172A
    COLOR_BLUE = RGBColor(37, 99, 235)      # #2563EB
    COLOR_DARK_BLUE = RGBColor(30, 58, 138) # #1E3A8A
    COLOR_BODY = RGBColor(51, 65, 85)       # #334155
    COLOR_MUTED = RGBColor(100, 116, 139)   # #64748B
    COLOR_PANEL = RGBColor(248, 250, 252)   # #F8FAFC
    COLOR_PANEL_ALT = RGBColor(241, 245, 249) # #F1F5F9
    COLOR_BORDER = RGBColor(203, 213, 225)  # #CBD5E1
    COLOR_LIGHT_BLUE = RGBColor(239, 246, 255) # #EFF6FF
    COLOR_BORDER_BLUE = RGBColor(191, 219, 254) # #BFDBFE

    # Source Dashboard Image
    src_img = r"C:\Users\Lenovo\.gemini\antigravity-ide\brain\917e1eb6-a0eb-4f5c-8ef8-9f948faa565d\.user_uploaded\media_1790002296382.png"
    dest_img = r"d:\projects\soc-project\presentation\dashboard_preview.png"
    if os.path.exists(src_img):
        shutil.copyfile(src_img, dest_img)

    def add_header(slide, title_text, category_text, slide_num):
        # Top line accent
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = COLOR_BLUE
        top_bar.line.color.rgb = COLOR_BLUE

        # Header branding
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.4))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = "SHIELDONE  |  SECURITY OPERATIONS CENTER PLATFORM"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_BLUE
        p.font.name = "Arial"

        # Slide Number
        slide_num_box = slide.shapes.add_textbox(Inches(10.5), Inches(0.4), Inches(2.0), Inches(0.4))
        tf_num = slide_num_box.text_frame
        tf_num.word_wrap = True
        tf_num.margin_left = tf_num.margin_top = tf_num.margin_right = tf_num.margin_bottom = 0
        p_num = tf_num.paragraphs[0]
        p_num.text = f"Slide {slide_num} of 7"
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.font.size = Pt(9.5)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_MUTED
        p_num.font.name = "Arial"

        # Kicker & Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.733), Inches(0.9))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        
        p_kicker = tf_t.paragraphs[0]
        p_kicker.text = category_text.upper()
        p_kicker.font.size = Pt(8.5)
        p_kicker.font.bold = True
        p_kicker.font.color.rgb = COLOR_BLUE
        p_kicker.font.name = "Arial"

        p_main = tf_t.add_paragraph()
        p_main.text = title_text
        p_main.font.size = Pt(22)
        p_main.font.bold = True
        p_main.font.color.rgb = COLOR_NAVY
        p_main.font.name = "Arial"

        # Blue underline
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.68), Inches(0.6), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_BLUE
        line.line.color.rgb = COLOR_BLUE

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    
    # Top bar
    top_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_BLUE
    top_bar.line.color.rgb = COLOR_BLUE

    # Category Pill
    cat_shape = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.6), Inches(1.5), Inches(4.133), Inches(0.45))
    cat_shape.fill.solid()
    cat_shape.fill.fore_color.rgb = COLOR_PANEL_ALT
    cat_shape.line.color.rgb = COLOR_BORDER
    tf_cat = cat_shape.text_frame
    tf_cat.word_wrap = True
    p_c = tf_cat.paragraphs[0]
    p_c.text = "ACADEMIC PROJECT PRESENTATION"
    p_c.alignment = PP_ALIGN.CENTER
    p_c.font.size = Pt(10)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_NAVY
    p_c.font.name = "Arial"

    # Main Title & Subtitle Box
    tbox = slide1.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.333), Inches(2.2))
    tf1 = tbox.text_frame
    tf1.word_wrap = True
    
    p1 = tf1.paragraphs[0]
    p1.text = "ShieldOne"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.size = Pt(48)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_NAVY
    p1.font.name = "Arial"

    p2 = tf1.add_paragraph()
    p2.text = "Security Operations Center — Real-Time Threat Monitoring Platform"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(16)
    p2.font.color.rgb = COLOR_BODY
    p2.font.name = "Arial"
    p2.space_before = Pt(8)

    # Presenter Card
    card1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.2), Inches(4.6), Inches(6.933), Inches(1.3))
    card1.fill.solid()
    card1.fill.fore_color.rgb = COLOR_PANEL
    card1.line.color.rgb = COLOR_BORDER

    # Text inside Presenter Card
    p_box = slide1.shapes.add_textbox(Inches(3.4), Inches(4.75), Inches(6.533), Inches(1.0))
    tf_p = p_box.text_frame
    tf_p.word_wrap = True

    # 3 Column text layout
    p_head = tf_p.paragraphs[0]
    p_head.text = "PRESENTED BY                     UID                                SECTION"
    p_head.font.size = Pt(9.5)
    p_head.font.bold = True
    p_head.font.color.rgb = COLOR_MUTED
    p_head.font.name = "Arial"

    p_val = tf_p.add_paragraph()
    p_val.text = "Bhargav                                24BCS11028                  803A"
    p_val.font.size = Pt(14)
    p_val.font.bold = True
    p_val.font.color.rgb = COLOR_NAVY
    p_val.font.name = "Arial"
    p_val.space_before = Pt(6)

    # =========================================================================
    # SLIDE 2: PROBLEM STATEMENT & PROPOSED SOLUTION
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Problem Statement & Proposed Solution", "Motivation & Overview", 2)

    # Left Box: Problem
    left_card = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.8))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = COLOR_PANEL
    left_card.line.color.rgb = COLOR_BORDER

    tb_left = slide2.shapes.add_textbox(Inches(1.05), Inches(2.2), Inches(5.1), Inches(4.4))
    tf_l = tb_left.text_frame
    tf_l.word_wrap = True
    
    pl_title = tf_l.paragraphs[0]
    pl_title.text = "CURRENT INDUSTRY CHALLENGES"
    pl_title.font.size = Pt(11)
    pl_title.font.bold = True
    pl_title.font.color.rgb = COLOR_NAVY
    pl_title.font.name = "Arial"
    pl_title.space_after = Pt(12)

    problems = [
        ("Continuous Monitoring:", "Cyber threats operate 24/7, requiring real-time streaming detection rather than periodic log reviews."),
        ("Fragmented Feeds:", "Threat intelligence is dispersed across disparate, uncoordinated external repositories."),
        ("Prioritization Bottleneck:", "Raw indicator logs lack context, geolocation mapping, and severity scoring."),
        ("Analyst Fatigue:", "Security teams need rapid threat interpretation to reduce MTTD and MTTR.")
    ]
    for bold_txt, norm_txt in problems:
        p_item = tf_l.add_paragraph()
        p_item.space_after = Pt(10)
        p_item.font.size = Pt(11)
        p_item.font.name = "Arial"
        
        run_b = p_item.add_run()
        run_b.text = "•  " + bold_txt + " "
        run_b.font.bold = True
        run_b.font.color.rgb = COLOR_NAVY
        
        run_n = p_item.add_run()
        run_n.text = norm_txt
        run_n.font.color.rgb = COLOR_BODY

    # Right Box: Solution
    right_card = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(2.0), Inches(5.7), Inches(4.8))
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = COLOR_PANEL
    right_card.line.color.rgb = COLOR_BORDER

    tb_right = slide2.shapes.add_textbox(Inches(7.05), Inches(2.2), Inches(5.2), Inches(4.4))
    tf_r = tb_right.text_frame
    tf_r.word_wrap = True
    
    pr_title = tf_r.paragraphs[0]
    pr_title.text = "THE SHIELDONE PLATFORM"
    pr_title.font.size = Pt(11)
    pr_title.font.bold = True
    pr_title.font.color.rgb = COLOR_BLUE
    pr_title.font.name = "Arial"
    pr_title.space_after = Pt(12)

    solutions = [
        ("Unified Ingestion:", "Collects multi-source threat intelligence feeds in real time."),
        ("Global Visualization:", "Maps active cyberattacks and geographical distributions on a live map."),
        ("Incident Escalation:", "Routes critical security events to analyst queues with SLA metrics."),
        ("AI Explanation:", "Utilizes LLM intelligence for instant threat context and mitigation guidance.")
    ]
    for bold_txt, norm_txt in solutions:
        p_item = tf_r.add_paragraph()
        p_item.space_after = Pt(10)
        p_item.font.size = Pt(11)
        p_item.font.name = "Arial"
        
        run_b = p_item.add_run()
        run_b.text = "•  " + bold_txt + " "
        run_b.font.bold = True
        run_b.font.color.rgb = COLOR_BLUE
        
        run_n = p_item.add_run()
        run_n.text = norm_txt
        run_n.font.color.rgb = COLOR_BODY

    # =========================================================================
    # SLIDE 3: WHAT SHIELDONE DOES (5 MODULES)
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "What ShieldOne Does", "Core Capabilities", 3)

    modules_data = [
        ("MODULE 01", "Dashboard", ["Live attack stream", "Interactive world map", "Analytics charts", "Dynamic risk gauge"], "REAL-TIME TELEMETRY"),
        ("MODULE 02", "Incidents", ["Escalated threats", "Analyst assignment", "MTTD / MTTR metrics", "Lifecycle queue"], "TRIAGE & SLA"),
        ("MODULE 03", "Threat Intel", ["Actively exploited CVEs", "Ransomware tracking", "CISA advisories", "Catalog inspection"], "CISA & ABUSE.CH"),
        ("MODULE 04", "Endpoints", ["Workstation health", "Server status checks", "Vulnerability tallies", "OS classification"], "HOST VISIBILITY"),
        ("MODULE 05", "Compliance", ["SOC 2 Type II", "ISO 27001 audit", "GDPR verification", "PCI DSS compliance"], "AUDIT FRAMEWORKS")
    ]

    card_width = 2.2
    gap = 0.18
    start_x = 0.8
    for i, (mod_num, title, bullets, tag) in enumerate(modules_data):
        cx = start_x + i * (card_width + gap)
        mcard = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(cx), Inches(2.0), Inches(card_width), Inches(4.8))
        mcard.fill.solid()
        mcard.fill.fore_color.rgb = COLOR_PANEL
        mcard.line.color.rgb = COLOR_BORDER

        tb_m = slide3.shapes.add_textbox(Inches(cx + 0.12), Inches(2.15), Inches(card_width - 0.24), Inches(4.5))
        tf_m = tb_m.text_frame
        tf_m.word_wrap = True

        p_num = tf_m.paragraphs[0]
        p_num.text = mod_num
        p_num.font.size = Pt(8.5)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_BLUE
        p_num.font.name = "Arial"

        p_t = tf_m.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_NAVY
        p_t.font.name = "Arial"
        p_t.space_after = Pt(12)

        for b in bullets:
            p_b = tf_m.add_paragraph()
            p_b.text = f"—  {b}"
            p_b.font.size = Pt(10)
            p_b.font.color.rgb = COLOR_BODY
            p_b.font.name = "Arial"
            p_b.space_after = Pt(6)

        p_tag = tf_m.add_paragraph()
        p_tag.text = tag
        p_tag.font.size = Pt(8)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_MUTED
        p_tag.font.name = "Arial"
        p_tag.space_before = Pt(24)

    # =========================================================================
    # SLIDE 4: REAL-TIME THREAT INTELLIGENCE
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Real-Time Threat Intelligence Sources", "Data Ingestion", 4)

    # Left Feeds
    left_feeds = [
        ("Feodo Tracker", "Abuse.ch", "Botnet Command & Control (C2) IP blocklists"),
        ("MalwareBazaar", "Abuse.ch", "Malware samples, payload signatures & family classifications")
    ]
    for i, (fname, org, desc) in enumerate(left_feeds):
        fy = 2.2 + i * 2.2
        fcard = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(fy), Inches(4.6), Inches(1.8))
        fcard.fill.solid()
        fcard.fill.fore_color.rgb = COLOR_PANEL
        fcard.line.color.rgb = COLOR_BORDER

        tb_f = slide4.shapes.add_textbox(Inches(1.0), Inches(fy + 0.2), Inches(4.2), Inches(1.4))
        tf_f = tb_f.text_frame
        tf_f.word_wrap = True

        p_fn = tf_f.paragraphs[0]
        p_fn.text = f"{fname}  [{org}]"
        p_fn.font.size = Pt(12.5)
        p_fn.font.bold = True
        p_fn.font.color.rgb = COLOR_NAVY
        p_fn.font.name = "Arial"

        p_fd = tf_f.add_paragraph()
        p_fd.text = desc
        p_fd.font.size = Pt(10.5)
        p_fd.font.color.rgb = COLOR_BODY
        p_fd.font.name = "Arial"
        p_fd.space_before = Pt(6)

    # Central Hub
    hub_card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.7), Inches(2.6), Inches(2.2), Inches(2.8))
    hub_card.fill.solid()
    hub_card.fill.fore_color.rgb = COLOR_LIGHT_BLUE
    hub_card.line.color.rgb = COLOR_BLUE

    tb_hub = slide4.shapes.add_textbox(Inches(5.8), Inches(3.2), Inches(2.0), Inches(2.0))
    tf_h = tb_hub.text_frame
    tf_h.word_wrap = True
    
    ph1 = tf_h.paragraphs[0]
    ph1.text = "ShieldOne"
    ph1.alignment = PP_ALIGN.CENTER
    ph1.font.size = Pt(16)
    ph1.font.bold = True
    ph1.font.color.rgb = COLOR_NAVY
    ph1.font.name = "Arial"

    ph2 = tf_h.add_paragraph()
    ph2.text = "Ingestion Engine"
    ph2.alignment = PP_ALIGN.CENTER
    ph2.font.size = Pt(9.5)
    ph2.font.bold = True
    ph2.font.color.rgb = COLOR_BLUE
    ph2.font.name = "Arial"

    ph3 = tf_h.add_paragraph()
    ph3.text = "30-Min Feed Sync"
    ph3.alignment = PP_ALIGN.CENTER
    ph3.font.size = Pt(8.5)
    ph3.font.color.rgb = COLOR_MUTED
    ph3.font.name = "Arial"
    ph3.space_before = Pt(10)

    # Right Feeds
    right_feeds = [
        ("ThreatFox", "Abuse.ch", "Indicators of Compromise (IOCs) with IPs & ports"),
        ("CISA KEV", "Official Feed", "Catalog of Known Exploited Vulnerabilities"),
        ("DShield / SANS ISC", "Sensor Net", "Honeypot-derived top hostile attacking IP addresses")
    ]
    for i, (fname, org, desc) in enumerate(right_feeds):
        fy = 2.0 + i * 1.55
        fcard = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.2), Inches(fy), Inches(4.3), Inches(1.35))
        fcard.fill.solid()
        fcard.fill.fore_color.rgb = COLOR_PANEL
        fcard.line.color.rgb = COLOR_BORDER

        tb_f = slide4.shapes.add_textbox(Inches(8.4), Inches(fy + 0.15), Inches(3.9), Inches(1.1))
        tf_f = tb_f.text_frame
        tf_f.word_wrap = True

        p_fn = tf_f.paragraphs[0]
        p_fn.text = f"{fname}  [{org}]"
        p_fn.font.size = Pt(11.5)
        p_fn.font.bold = True
        p_fn.font.color.rgb = COLOR_NAVY
        p_fn.font.name = "Arial"

        p_fd = tf_f.add_paragraph()
        p_fd.text = desc
        p_fd.font.size = Pt(9.5)
        p_fd.font.color.rgb = COLOR_BODY
        p_fd.font.name = "Arial"
        p_fd.space_before = Pt(3)

    # =========================================================================
    # SLIDE 5: HOW SHIELDONE WORKS (ARCHITECTURE PIPELINE)
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "How ShieldOne Works", "Architecture & Data Flow", 5)

    arch_steps = [
        ("Step 1", "Threat Feeds", "5 Open external intelligence sources"),
        ("Step 2", "Backend API", "Node.js + Express ingestion server"),
        ("Step 3", "Processing", "GeoIP lookup & severity computation"),
        ("Step 4", "Socket.io", "WebSocket live streaming pipeline"),
        ("Step 5", "Dashboard", "Feed, Map, Charts & Risk Gauge")
    ]
    swidth = 2.2
    sgap = 0.18
    sstart = 0.8
    for i, (st, name, desc) in enumerate(arch_steps):
        sx = sstart + i * (swidth + sgap)
        scard = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(sx), Inches(2.2), Inches(swidth), Inches(3.0))
        scard.fill.solid()
        if i == 3: # Highlight Socket.io
            scard.fill.fore_color.rgb = COLOR_LIGHT_BLUE
            scard.line.color.rgb = COLOR_BLUE
        else:
            scard.fill.fore_color.rgb = COLOR_PANEL
            scard.line.color.rgb = COLOR_BORDER

        tb_s = slide5.shapes.add_textbox(Inches(sx + 0.15), Inches(2.4), Inches(swidth - 0.3), Inches(2.5))
        tf_s = tb_s.text_frame
        tf_s.word_wrap = True

        p_st = tf_s.paragraphs[0]
        p_st.text = st
        p_st.font.size = Pt(9)
        p_st.font.bold = True
        p_st.font.color.rgb = COLOR_BLUE
        p_st.font.name = "Arial"

        p_sn = tf_s.add_paragraph()
        p_sn.text = name
        p_sn.font.size = Pt(13)
        p_sn.font.bold = True
        p_sn.font.color.rgb = COLOR_NAVY
        p_sn.font.name = "Arial"
        p_sn.space_before = Pt(4)

        p_sd = tf_s.add_paragraph()
        p_sd.text = desc
        p_sd.font.size = Pt(10)
        p_sd.font.color.rgb = COLOR_BODY
        p_sd.font.name = "Arial"
        p_sd.space_before = Pt(10)

    # Bottom Callout Box
    callout = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(5.6), Inches(11.733), Inches(1.1))
    callout.fill.solid()
    callout.fill.fore_color.rgb = COLOR_PANEL
    callout.line.color.rgb = COLOR_BLUE

    tb_c = slide5.shapes.add_textbox(Inches(1.1), Inches(5.8), Inches(11.133), Inches(0.7))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True
    
    pc1 = tf_c.paragraphs[0]
    r1 = pc1.add_run()
    r1.text = "Real-Time Delivery Mechanism: "
    r1.font.bold = True
    r1.font.size = Pt(12)
    r1.font.color.rgb = COLOR_NAVY
    
    r2 = pc1.add_run()
    r2.text = "Attack events are pushed to connected browser clients every 2 seconds using WebSockets without requiring page refresh (Zero-Polling Architecture)."
    r2.font.size = Pt(12)
    r2.font.color.rgb = COLOR_BODY

    # =========================================================================
    # SLIDE 6: AI FEATURES & LIVE DASHBOARD INTERFACE
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "AI Engine & Live Dashboard Interface", "Intelligence & Interface", 6)

    # Left Column: AI Features
    ai_card = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.0), Inches(5.0), Inches(4.8))
    ai_card.fill.solid()
    ai_card.fill.fore_color.rgb = COLOR_PANEL
    ai_card.line.color.rgb = COLOR_BORDER

    tb_ai = slide6.shapes.add_textbox(Inches(1.05), Inches(2.2), Inches(4.5), Inches(4.4))
    tf_ai = tb_ai.text_frame
    tf_ai.word_wrap = True

    p_ait = tf_ai.paragraphs[0]
    p_ait.text = "AI FEATURES & TRIAGE"
    p_ait.font.size = Pt(11)
    p_ait.font.bold = True
    p_ait.font.color.rgb = COLOR_NAVY
    p_ait.font.name = "Arial"
    p_ait.space_after = Pt(12)

    ai_items = [
        ("AI Attack Explainer:", "1-click contextual breakdowns for incoming security alerts."),
        ("SOC Chatbot Assistant:", "Interactive natural language assistant for threat investigation."),
        ("Underlying LLM:", "Google Gemini 2.5 Flash through OpenRouter API."),
        ("Mitigation Guidance:", "Root-cause attack pattern analysis and step-by-step mitigation guidance.")
    ]
    for bold_txt, norm_txt in ai_items:
        p_item = tf_ai.add_paragraph()
        p_item.space_after = Pt(10)
        p_item.font.size = Pt(11)
        p_item.font.name = "Arial"
        
        run_b = p_item.add_run()
        run_b.text = "•  " + bold_txt + " "
        run_b.font.bold = True
        run_b.font.color.rgb = COLOR_NAVY
        
        run_n = p_item.add_run()
        run_n.text = norm_txt
        run_n.font.color.rgb = COLOR_BODY

    # Right Column: Dashboard Screenshot
    if os.path.exists(dest_img):
        dash_pic = slide6.shapes.add_picture(dest_img, Inches(6.1), Inches(2.0), width=Inches(6.433))
        # Add border around screenshot
        dash_pic_border = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.08), Inches(1.98), Inches(6.47), Inches(4.84))
        dash_pic_border.fill.background()
        dash_pic_border.line.color.rgb = COLOR_BORDER

    # =========================================================================
    # SLIDE 7: TECHNOLOGY STACK & CONCLUSION
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Technology Stack & Conclusion", "Summary", 7)

    # Left: Tech Stack Rows
    tech_categories = [
        ("FRONTEND STACK", "React 19, Vite, Leaflet, React-Leaflet, Recharts, GSAP Animations, Socket.io-client"),
        ("BACKEND & INFRASTRUCTURE", "Node.js, Express 5, Socket.io (WebSockets), Axios, Supabase, Nodemailer"),
        ("INTELLIGENCE & AI ENGINE", "Feodo Tracker, MalwareBazaar, ThreatFox, CISA KEV, DShield / SANS, Gemini 2.5 Flash")
    ]
    for i, (cat, items) in enumerate(tech_categories):
        ty = 2.0 + i * 1.55
        tcard = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(ty), Inches(6.2), Inches(1.35))
        tcard.fill.solid()
        tcard.fill.fore_color.rgb = COLOR_PANEL
        tcard.line.color.rgb = COLOR_BORDER

        tb_t = slide7.shapes.add_textbox(Inches(1.05), Inches(ty + 0.15), Inches(5.7), Inches(1.05))
        tf_t = tb_t.text_frame
        tf_t.word_wrap = True

        p_c = tf_t.paragraphs[0]
        p_c.text = cat
        p_c.font.size = Pt(9)
        p_c.font.bold = True
        p_c.font.color.rgb = COLOR_BLUE
        p_c.font.name = "Arial"

        p_i = tf_t.add_paragraph()
        p_i.text = items
        p_i.font.size = Pt(11)
        p_i.font.bold = True
        p_i.font.color.rgb = COLOR_NAVY
        p_i.font.name = "Arial"
        p_i.space_before = Pt(4)

    # Right: Formal Conclusion Card
    concl_card = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.4), Inches(2.0), Inches(5.133), Inches(4.8))
    concl_card.fill.solid()
    concl_card.fill.fore_color.rgb = COLOR_PANEL
    concl_card.line.color.rgb = COLOR_NAVY

    tb_concl = slide7.shapes.add_textbox(Inches(7.7), Inches(2.6), Inches(4.533), Inches(3.6))
    tf_concl = tb_concl.text_frame
    tf_concl.word_wrap = True

    p_cq = tf_concl.paragraphs[0]
    p_cq.text = "“ShieldOne — From raw threat intelligence to real-time security awareness.”"
    p_cq.alignment = PP_ALIGN.CENTER
    p_cq.font.size = Pt(18)
    p_cq.font.bold = True
    p_cq.font.color.rgb = COLOR_NAVY
    p_cq.font.name = "Arial"

    p_cs = tf_concl.add_paragraph()
    p_cs.text = "Unified threat visibility, rapid automated triage, and proactive incident response for modern security operations."
    p_cs.alignment = PP_ALIGN.CENTER
    p_cs.font.size = Pt(11)
    p_cs.font.color.rgb = COLOR_BODY
    p_cs.font.name = "Arial"
    p_cs.space_before = Pt(16)

    # Save Output Presentation
    output_path = r"d:\projects\soc-project\presentation\ShieldOne_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    create_shieldone_presentation()
